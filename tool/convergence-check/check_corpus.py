"""Exportación reproducible y comprobación del registro, CSV y SoA publicados."""
from pathlib import Path
from collections import Counter
import argparse
import csv
import io
import json
import re
import sys
import openpyxl
from project import cargar, clasificacion, ESCALAS

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ENTREGA = ROOT / "deliverables/fase1/revision_2026-09-27"
REG = ENTREGA / "SGSIREG001002_Registro_y_Plan_Riesgos_v2_0.xlsx"
SOA = ENTREGA / "SGSIREG003_Declaracion_Aplicabilidad_v1_10.xlsx"
RFIELDS = ["codigo", "titulo", "dominios", *ESCALAS, "propietario", "justificacion_dominios"]
CFIELDS = ["id_control", "id_riesgo", "descripcion", "dominio", "estado", "evidencia", "periodicidad_dias", "fecha_ultima_verificacion", "alcance_cobertura", "condicion"]
SIMULADOS = {"estado", "evidencia", "fecha_ultima_verificacion"}
ORDEN = {"PLANIFICADO": 0, "PARCIAL": 1, "IMPLEMENTADO": 2}


def extraer():
    wb = openpyxl.load_workbook(REG, data_only=True)
    riesgos, controles = [], []
    cadenas = {}
    for row in wb["Registro de riesgos"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        valores = [row[0], row[1], row[4], *row[6:10], row[11], row[14], row[17]]
        riesgo = dict(zip(RFIELDS, valores))
        esperado = clasificacion(*(riesgo[k] for k in ESCALAS))
        if row[13] != esperado or row[10] != max(row[6:10]) or row[12] != max(row[6:10])*row[11]:
            raise ValueError(f"{row[0]}: fórmulas o resultados del Excel no coinciden con la metodología; recalcular el libro")
        riesgos.append(riesgo)
        cadenas[row[0]] = len(row[2].split("→"))
    for row in wb["Controles"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            eslabon = re.search(r"Eslabón (\d+)", row[3] or "")
            if row[1] not in cadenas or not row[3]:
                raise ValueError(f"{row[0]}: riesgo o eslabón sin traza")
            if eslabon and not 1 <= int(eslabon[1]) <= cadenas[row[1]]:
                raise ValueError(f"{row[0]}: eslabón fuera de la cadena resumida")
            controles.append(dict(zip(CFIELDS, [row[0],row[1],row[2],row[4],row[5],row[6] or "",row[9],row[10] or "",row[12],row[13] or ""])))
    return riesgos, controles


def texto_csv(filas, campos):
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=campos, lineterminator="\n")
    writer.writeheader()
    writer.writerows(filas)
    return out.getvalue()


def comprobar_soa(controles):
    wb = openpyxl.load_workbook(SOA,data_only=True)
    filas = list(wb["SoA"].iter_rows(min_row=2,values_only=True))
    esperados = {f"{bloque}.{n}" for bloque,cantidad in [(5,37),(6,8),(7,14),(8,34)] for n in range(1,cantidad+1)}
    codigos = [str(r[0]) for r in filas if r[0]]
    if len(codigos)!=93 or set(codigos)!=esperados:
        raise ValueError("SoA: los 93 identificadores del Anexo A no son únicos/completos")
    ids = {c["id_control"]:c for c in controles}
    excepciones = []
    clases = Counter()
    for row in filas:
        codigo,aplicable,estado,traza = str(row[0]),row[2],row[6],row[7] or ""
        if aplicable not in {"Sí","No"} or (aplicable=="No") != (estado=="NO APLICA"):
            raise ValueError(f"SoA {codigo}: aplicabilidad/estado incoherentes")
        if aplicable=="Sí" and estado not in ORDEN:
            raise ValueError(f"SoA {codigo}: estado desconocido")
        clase = ("NO APLICA" if aplicable=="No" else
                 "Registro" if traza.startswith("C-") else
                 "Documento" if traza.startswith(("DOC-","REG-","Contrato")) else
                 "Deuda" if traza.startswith("Sin soporte") else
                 "Otra" if traza else "Sin traza")
        if len(row)<9 or row[8]!=clase:
            raise ValueError(f"SoA {codigo}: clase de traza calculada incorrecta")
        clases[clase] += 1
        # No confundir RSC-017 con C-017.
        refs=set(re.findall(r"(?<![A-Z])C-\d{3}",traza))
        if refs-ids.keys():
            raise ValueError(f"SoA {codigo}: referencias inexistentes {sorted(refs-ids.keys())}")
        if refs:
            minimo=min((ids[x]["estado"] for x in refs),key=ORDEN.get)
            if estado!=minimo:
                if codigo=="7.1" and estado=="PARCIAL" and minimo=="PLANIFICADO":
                    excepciones.append("A.7.1: perímetro físico preexistente, PARCIAL, excepción de DOC-001 §9")
                else:
                    raise ValueError(f"SoA {codigo}: {estado}, registro {minimo}")
        elif codigo=="8.14":
            if estado!="IMPLEMENTADO" or "Contrato" not in traza:
                raise ValueError("SoA A.8.14: excepción contractual alterada; revisar")
            excepciones.append("A.8.14: contrato de tercero, sin control propio del registro")
        elif aplicable=="Sí" and not traza:
            raise ValueError(f"SoA {codigo}: falta traza; declarar deuda si procede")
    aplicables = sum(r[2]=="Sí" for r in filas)
    resumen = {3:len(filas),4:aplicables,5:clases['NO APLICA'],
               **{n:sum(r[6]==estado for r in filas) for n,estado in [(6,'IMPLEMENTADO'),(7,'PARCIAL'),(8,'PLANIFICADO')]},
               **{n:clases[clase] for n,clase in [(11,'Registro'),(12,'Documento'),(13,'Deuda'),(14,'Otra'),(15,'Sin traza')]},
               16:sum(clases[k] for k in ['Registro','Documento','Deuda','Sin traza'])}
    for n,esperado in resumen.items():
        if wb['Resumen'][f'B{n}'].value!=esperado:
            raise ValueError(f"SoA Resumen B{n}: resultado calculado incorrecto")
    if clases['Otra'] or clases['Sin traza'] or resumen[16]!=aplicables:
        raise ValueError("SoA: trazas pendientes de clasificar o cuadre incompleto")
    return excepciones


def ejecutar(exportar=False):
    riesgos, controles = extraer()
    excepciones = comprobar_soa(controles)
    for filename,rows,fields in [("riesgos.csv",riesgos,RFIELDS),("controles.csv",controles,CFIELDS)]:
        esperado=texto_csv(rows,fields)
        path=HERE/filename
        if exportar:
            path.write_text(esperado,encoding="utf-8",newline="")
        elif path.read_text(encoding="utf-8-sig")!=esperado:
            raise ValueError(f"{filename}: difiere del Excel; exportar y revisar")
    cargar(HERE/"riesgos.csv",HERE/"controles.csv")
    _, demo=cargar(HERE/"riesgos.csv",HERE/"controles_junio.csv")
    source={c['id_control']:c for c in controles}
    if {c['id_control'] for c in demo}!=set(source):
        raise ValueError("Demostración: identificadores diferentes del registro")
    for c in demo:
        for k in set(CFIELDS)-SIMULADOS:
            if c[k]!=source[c['id_control']][k]:
                raise ValueError(f"Demostración {c['id_control']}: difiere {k}")
    return {"riesgos":len(riesgos),"controles":len(controles),"soa":93,"excepciones":excepciones,
            "limite":"Comprueba datos y estados declarados; no verifica evidencias ni cierra HD-01/HD-06 por sí solo."}


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--exportar",action="store_true",help="Actualizar los CSV de corte desde el Excel; la demostración se revisa aparte")
    args=parser.parse_args()
    try:
        print(json.dumps(ejecutar(args.exportar),ensure_ascii=False,indent=2))
    except (OSError,ValueError,KeyError) as exc:
        print(f"Error de coherencia: {exc}",file=sys.stderr)
        sys.exit(2)
